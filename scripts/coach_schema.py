"""Leest het trainingsschema van coach Kaj uit zijn Google Sheet.

De Sheet is de bron, niet deze code. Voegt Kaj een week toe, dan verschijnt die
de volgende ochtend vanzelf in Hevy zonder dat hier iets aan hoeft te veranderen.

Structuur van de Sheet (tabblad per blok):

    A2   WEEK 1
    A8   Maandag | September 07        E8  Primaire Dip + Secundaire Pull-Up
    A9   OEFENING | SETS | REPS | RPE/LOAD KG | OPMERKINGEN | WERKELIJKE RPE | Set 1..5
    A10  Competition Dip | 3 | 3 | 6
    ...  lege regel sluit de dag af

Schrijft data/coachschema.json weg.

EEN VALKUIL IN DE BRON
Google Sheets maakt van een repbereik als "6-12" een datum: 6 december. Deze
parser draait dat terug (dag-maand), maar het blijft een bron van fouten in de
Sheet zelf. Kaj kan dat voorkomen door die cellen als tekst op te maken.
"""

import datetime as dt
import io
import json
import pathlib
import re
import urllib.request

SHEET_ID = "1QGxPFQB17u9zDPNXMl74pcVUERlX2g7RwZBRk-Trpfo"
EXPORT = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"

DAGNAMEN = {"maandag": 0, "dinsdag": 1, "woensdag": 2, "donderdag": 3,
            "vrijdag": 4, "zaterdag": 5, "zondag": 6}

MAANDEN = {
    "january": 1, "januari": 1, "february": 2, "februari": 2, "march": 3, "maart": 3,
    "april": 4, "may": 5, "mei": 5, "june": 6, "juni": 6, "july": 7, "juli": 7,
    "august": 8, "augustus": 8, "september": 9, "october": 10, "oktober": 10,
    "november": 11, "december": 12,
}


def lees_datum(label):
    """'September 12' wordt een echte datum.

    Het jaartal staat er niet bij, dus we kiezen het jaar waarin die dag het
    dichtst bij vandaag ligt. Zonder datum kan een sessie aan de verkeerde dag
    gekoppeld worden, en dan komt een training in een blok terecht dat nog
    moet plaatsvinden.
    """
    if not label:
        return None
    m = re.search(r"([A-Za-z]+)\s+(\d{1,2})", label.strip())
    if not m:
        return None
    maand = MAANDEN.get(m.group(1).lower())
    if not maand:
        return None
    dag = int(m.group(2))
    vandaag = dt.date.today()
    kandidaten = []
    for jaar in (vandaag.year - 1, vandaag.year, vandaag.year + 1):
        try:
            kandidaten.append(dt.date(jaar, maand, dag))
        except ValueError:
            pass
    return min(kandidaten, key=lambda d: abs((d - vandaag).days)) if kandidaten else None


def ontdatum(waarde):
    """Draait de datumconversie van Google Sheets terug naar een repbereik."""
    if isinstance(waarde, (dt.datetime, dt.date)):
        return f"{waarde.day}-{waarde.month}"
    return waarde


def tekst(waarde):
    if waarde is None:
        return ""
    waarde = ontdatum(waarde)
    if isinstance(waarde, float) and waarde.is_integer():
        return str(int(waarde))
    return str(waarde).strip()


def haal_sheet():
    with urllib.request.urlopen(EXPORT, timeout=90) as r:
        return io.BytesIO(r.read())


def parse(bestand):
    import openpyxl
    wb = openpyxl.load_workbook(bestand, data_only=True)
    blokken = []
    for naam in wb.sheetnames:
        ws = wb[naam]
        week = None
        dagen, huidig = [], None
        for r in range(1, ws.max_row + 1):
            a = tekst(ws.cell(row=r, column=1).value)
            if not a:
                huidig = None                     # lege regel sluit de dag af
                continue
            if re.fullmatch(r"WEEK\s*\d+", a, re.I):
                week = int(re.search(r"\d+", a).group())
                continue
            if a.upper() == "OEFENING":
                continue                          # kolomkoppen overslaan
            dag = a.split("|")[0].strip().lower()
            if dag in DAGNAMEN:
                label = a.split("|")[-1].strip() if "|" in a else ""
                datum = lees_datum(label)
                huidig = {"dag": dag, "weekdag": DAGNAMEN[dag],
                          "datum_label": label,
                          "datum": datum.isoformat() if datum else None,
                          "thema": tekst(ws.cell(row=r, column=5).value),
                          "tabblad": naam, "kop_rij": r, "oefeningen": []}
                dagen.append(huidig)
                continue
            if huidig is None:
                continue
            huidig["oefeningen"].append({
                "rij": r,                     # waar de schrijver de sets neerzet
                "naam": a,
                "sets": tekst(ws.cell(row=r, column=2).value),
                "reps": tekst(ws.cell(row=r, column=3).value),
                "rpe_load": tekst(ws.cell(row=r, column=4).value),
                "opmerking": tekst(ws.cell(row=r, column=5).value),
            })
        if dagen:
            blokken.append({"blok": naam, "week": week, "dagen": dagen})
    return blokken


def main():
    blokken = parse(haal_sheet())
    uit = pathlib.Path("data")
    uit.mkdir(exist_ok=True)
    (uit / "coachschema.json").write_text(
        json.dumps({"opgehaald": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"),
                    "bron": EXPORT, "blokken": blokken}, indent=1, ensure_ascii=False),
        encoding="utf-8")
    for b in blokken:
        print(f"{b['blok']} — week {b['week']}: {len(b['dagen'])} dagen")
        for d in b["dagen"]:
            rijen = [o["rij"] for o in d["oefeningen"]]
            bereik = f"rij {min(rijen)}-{max(rijen)}" if rijen else "leeg"
            print(f"   {d['dag']:10s} {d.get('datum') or '?':10s} "
                  f"{len(d['oefeningen'])} oefeningen ({bereik}) — {d['thema'][:40]}")


if __name__ == "__main__":
    main()
