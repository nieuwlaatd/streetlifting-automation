"""Bouwt het liftoverzicht voor coach Kaj als Excel-bestand.

Kaj wil in een oogopslag zien waar Dylan staat: per wedstrijdlift een geschat
maximum, en daaronder de sets waar die schatting op rust. Dit bestand is een
momentopname uit Hevy; het levende logboek is zijn eigen Google Sheet.

WAAROM ER FORMULES IN STAAN EN GEEN KALE GETALLEN
Het lichaamsgewicht bepaalt bij dip, pull-up en muscle-up de helft van de
belasting. Weegt Dylan volgende maand 79 kilo, dan verandert elke schatting
mee. Daarom staat het gewicht op EEN plek (Methode!B3) en rekent de rest zich
daaruit. Kaj kan dat getal zelf aanpassen en ziet meteen wat het doet.

De schatting per set staat in kolom I van Wedstrijdlifts. Het overzicht neemt
daar niet het maximum van maar de op een na hoogste: een enkele te
optimistisch gelogde set bepaalt anders in zijn eentje het beeld.

Gebruik:  python scripts/liftoverzicht.py [pad naar .xlsx]
Nodig:    HEVY_API_KEY
"""

import datetime as dt
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import hevy_sync

STANDAARD_PAD = (r"C:\Users\Dylan\OneDrive\Documenten\Workout streetlifting"
                 r"\Liftoverzicht-Dylan-voor-Kaj.xlsx")

# Hevy-oefening naar de lift zoals Kaj hem noemt. Een muscle-up met band is
# geen wedstrijdlift maar hoort er wel bij: zonder die regels lijkt het alsof
# er niet aan gewerkt wordt.
WEDSTRIJD = {
    "Chest Dip (Weighted)": "Dip",
    "Chest Dip": "Dip",
    "Triceps Dip (Weighted)": "Dip",
    "Pull Up (Weighted)": "Pull-up",
    "Chin Up (Weighted)": "Pull-up (chin)",
    "Squat (Barbell)": "Squat",
    "Muscle Up": "Muscle-up",
    "Band Assisted Muscle-Up": "Muscle-up (band)",
}
# Deze tellen mee voor het geschatte maximum in het overzicht. De chin-up en de
# muscle-up met band staan er bewust niet bij: het zijn andere lifts.
TELT_MEE = {"Dip": ["Dip"], "Pull-up": ["Pull-up"], "Squat": ["Squat"],
            "Muscle-up": ["Muscle-up"]}
LICHAAMSGEBONDEN = {"Dip", "Pull-up", "Pull-up (chin)", "Muscle-up", "Muscle-up (band)"}
DOEL = {"Dip": 105, "Pull-up": 80, "Squat": 160, "Muscle-up": 22}

BLAUW = PatternFill("solid", fgColor="1B4E9B")
KOP = Font(name="Arial", size=11, bold=True, color="FFFFFF")
TITEL = Font(name="Arial", size=14, bold=True)
SUB = Font(name="Arial", size=9, color="666666")
VET = Font(name="Arial", size=11, bold=True)
GEWOON = Font(name="Arial", size=10)
GETAL = Font(name="Arial", size=10, bold=True)
INVOER = Font(name="Arial", size=10, color="0000FF")


def kop_rij(ws, rij, koppen, breedtes):
    for i, (naam, breedte) in enumerate(zip(koppen, breedtes), start=1):
        c = ws.cell(row=rij, column=i, value=naam)
        c.font, c.fill = KOP, BLAUW
        c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = breedte


def zet(ws, rij, kolom, waarde, font=GEWOON, fmt=None):
    c = ws.cell(row=rij, column=kolom, value=waarde)
    c.font = font
    if fmt:
        c.number_format = fmt
    return c


def werksets(oefening):
    return [s for s in oefening.get("sets", [])
            if s.get("type") != "warmup" and s.get("reps")]


def bouw(workouts, bw, pad):
    wb = openpyxl.Workbook()

    # ---------- Wedstrijdlifts ----------
    ws = wb.active
    ws.title = "Wedstrijdlifts"
    zet(ws, 1, 1, "Alle werksets op de wedstrijdlifts, oudste eerst.", TITEL)
    zet(ws, 2, 1, "Kolom I rekent het geschatte 1RM per set uit, via de RPE in K en het "
                  "aandeel van het 1RM in L. Bij dip, pull-up en muscle-up telt het "
                  "lichaamsgewicht mee (Methode!B3).", SUB)
    kop_rij(ws, 4, ["Datum", "Lift", "Sessie", "Type", "Gewicht", "Reps", "RPE",
                    "Systeem", "Geschat 1RM", "Notitie", "RPE gebruikt", "Aandeel 1RM"],
            [12, 16, 20, 10, 10, 8, 8, 10, 13, 46, 12, 12])

    rij = 5
    per_lift = {}
    for w in sorted(workouts, key=lambda x: x["start_time"]):
        datum = w["start_time"][:10]
        for oef in w.get("exercises", []):
            lift = WEDSTRIJD.get(oef.get("title") or "")
            if not lift:
                continue
            gebonden = lift in LICHAAMSGEBONDEN
            for s in werksets(oef):
                # Een muscle-up zonder gelogd gewicht is een poging, geen lift.
                # Dylan schreef er zelf bij dat hij de techniek nog niet heeft
                # en meteen valt; daar een maximum uit afleiden zou Kaj een
                # getal geven dat nergens op slaat. De regel blijft wel staan,
                # zodat te zien is dat eraan gewerkt wordt.
                poging = lift.startswith("Muscle-up") and s.get("weight_kg") is None
                zet(ws, rij, 1, datum)
                zet(ws, rij, 2, lift)
                zet(ws, rij, 3, w.get("title") or "")
                zet(ws, rij, 4, "poging" if poging else (s.get("type") or "normal"))
                zet(ws, rij, 5, float(s.get("weight_kg") or 0), fmt="0.#")
                zet(ws, rij, 6, s.get("reps"))
                zet(ws, rij, 7, s.get("rpe"))
                zet(ws, rij, 8, "=E{0}+Methode!$B$3".format(rij) if gebonden
                    else "=E{0}".format(rij), fmt="0.0")
                korting = "IF(F{0}>=7,0.96,IF(F{0}=6,0.98,1))".format(rij)
                zet(ws, rij, 9, "=H{0}/L{0}*{1}".format(rij, korting)
                    + ("-Methode!$B$3" if gebonden else ""), GETAL, fmt="0.0")
                zet(ws, rij, 10, (oef.get("notes") or "")[:200])

                # De twee tussenstappen staan als kolom, niet weggestopt in een
                # formule. Kaj kan dan zien waar een schatting vandaan komt, en
                # een half punt RPE gaat niet meer mis: 6,5 zit tussen twee
                # kolommen van de tabel in en wordt daartussen ingeschat.
                zet(ws, rij, 11,
                    '=IF(G{0}<>"",MIN(MAX(G{0},6),10),IF(D{0}="failure",10,""))'.format(rij),
                    fmt="0.0")
                cel = ("INDEX(Methode!$B$21:$F$28,MATCH(MIN(MAX(F{0},1),8),"
                       "Methode!$A$21:$A$28,0),{1})")
                onder = cel.format(rij, "MIN(INT(K{0}),9)-5".format(rij))
                boven = cel.format(rij, "MIN(INT(K{0}),9)-4".format(rij))
                zet(ws, rij, 12,
                    '=IF(K{0}="",1/(1+MIN(F{0},8)/30),'
                    "{1}+({2}-{1})*(K{0}-MIN(INT(K{0}),9)))".format(rij, onder, boven),
                    fmt="0.000")
                # Dezelfde rekensom als in kolom I, maar in Python, zodat het
                # overzicht hieronder kan aanwijzen op welke set het rust.
                schatting = None if poging else hevy_sync.e1rm(s, bw, gebonden)
                per_lift.setdefault(lift, []).append((schatting, rij))
                rij += 1
    laatste = rij - 1
    ws.auto_filter.ref = "A4:L{0}".format(laatste)
    ws.freeze_panes = "A5"

    # ---------- Overzicht ----------
    ov = wb.create_sheet("Overzicht", 0)
    eerste_datum = min(w["start_time"][:10] for w in workouts)
    laatste_datum = max(w["start_time"][:10] for w in workouts)
    zet(ov, 1, 1, "Dylan \u2014 streetlifting, klasse \u221280 kg", TITEL)
    zet(ov, 2, 1, "Uitdraai voor coach Kaj \u00b7 gegevens uit Hevy tot en met "
                  + laatste_datum, SUB)
    zet(ov, 4, 1, "Atleet", VET)
    for i, (label, waarde) in enumerate([
            ("Lichaamsgewicht", "=Methode!B3"), ("Lengte", "186 cm"),
            ("Wedstrijdklasse", "\u221280 kg"), ("Doelwedstrijd", "circa september 2027"),
            ("Gelogde sessies", len(workouts)),
            ("Periode", eerste_datum + " t/m " + laatste_datum)], start=5):
        zet(ov, i, 1, label)
        zet(ov, i, 2, waarde)

    zet(ov, 12, 1, "Stand per wedstrijdlift", VET)
    kop_rij(ov, 13, ["Lift", "Geschat 1RM", "Maatgevende set", "Datum", "Doel wedstrijd",
                     "Nog te gaan", "Werksets gelogd"], [16, 14, 18, 12, 14, 12, 15])
    for i, lift in enumerate(["Dip", "Pull-up", "Squat", "Muscle-up"], start=14):
        rijen = [p for naam in TELT_MEE[lift] for p in per_lift.get(naam, [])
                 if p[0] is not None]
        zet(ov, i, 1, lift)
        if rijen:
            # De op een na hoogste van de drie beste, niet de allerhoogste: een
            # losse uitschieter of een set die te gretig als falen is gelogd
            # bepaalt anders in zijn eentje het beeld.
            top = sorted(rijen, reverse=True)[:3]
            _, bron = top[len(top) // 2] if len(top) >= 3 else top[-1]
            # Verwijzing naar die ene cel in plaats van een matrixformule: een
            # LARGE(IF(...)) rekent in Excel goed maar in Google Sheets stil
            # verkeerd, en een verkeerd getal is erger dan een omweg. Verandert
            # het lichaamsgewicht, dan loopt deze cel gewoon mee.
            zet(ov, i, 2, "=Wedstrijdlifts!I{0}".format(bron), GETAL, fmt="0.0")
            zet(ov, i, 3, "{0:g} kg \u00d7 {1}".format(ws.cell(row=bron, column=5).value,
                                                       ws.cell(row=bron, column=6).value))
            zet(ov, i, 4, ws.cell(row=bron, column=1).value)
        else:
            zet(ov, i, 2, "geen data")
            zet(ov, i, 3, "\u2014")
            zet(ov, i, 4, "\u2014")
        zet(ov, i, 5, DOEL[lift])
        zet(ov, i, 6, "=IF(ISNUMBER(B{0}),E{0}-B{0},E{0})".format(i), fmt="0.0")
        zet(ov, i, 7, len(rijen))
    zet(ov, 18, 1, "Totaal", VET)
    zet(ov, 18, 2, "=SUM(B14:B17)", GETAL, fmt="0.0")
    zet(ov, 18, 5, "=SUM(E14:E17)", GETAL, fmt="0.0")
    zet(ov, 18, 6, "=E18-B18", GETAL, fmt="0.0")
    ov.freeze_panes = "A14"

    zet(ov, 20, 1, "Toelichting", VET)
    for i, regel in enumerate([
            "Dip, pull-up en muscle-up zijn gescoord op toegevoegd gewicht; de schatting "
            "loopt over de systeembelasting (lichaamsgewicht plus schijf).",
            "Het totaal is de streetliftingsom van de vier lifts.",
            "Doel 367 kg is gebaseerd op de middenmoot van het DSN-veld 2025 in de "
            "\u221280 kg klasse.",
            "Grootste achterstanden: dip en muscle-up. Daar zit ook de meeste ruimte.",
            "De muscle-up staat op 'geen data': de pogingen tot nu toe zijn zonder "
            "gelogd gewicht en volgens Dylans eigen notitie technisch nog niet rond. "
            "Ze staan wel in het tabblad Wedstrijdlifts, gemarkeerd als poging.",
            "Sinds 7 september 2026 volgt Dylan het schema van Kaj; de uitvoering per "
            "sessie staat in Kajs eigen Google Sheet."], start=21):
        zet(ov, i, 1, regel)

    # ---------- Alle oefeningen ----------
    ao = wb.create_sheet("Alle oefeningen")
    zet(ao, 1, 1, "Elke oefening: volume en zwaarste set over de hele periode.", TITEL)
    kop_rij(ao, 3, ["Oefening", "Werksets", "Sessies", "Eerste", "Laatste",
                    "Zwaarste set", "Tot falen", "Met RPE"],
            [34, 10, 10, 12, 12, 16, 10, 10])
    stat = {}
    for w in workouts:
        datum = w["start_time"][:10]
        for oef in w.get("exercises", []):
            t = oef.get("title") or ""
            d = stat.setdefault(t, {"sets": 0, "dagen": set(), "eerste": datum,
                                    "laatste": datum, "zwaarst": (0, 0),
                                    "falen": 0, "rpe": 0})
            for s in werksets(oef):
                d["sets"] += 1
                d["dagen"].add(datum)
                d["eerste"] = min(d["eerste"], datum)
                d["laatste"] = max(d["laatste"], datum)
                paar = (float(s.get("weight_kg") or 0), s.get("reps") or 0)
                if paar > d["zwaarst"]:
                    d["zwaarst"] = paar
                d["falen"] += s.get("type") == "failure"
                d["rpe"] += s.get("rpe") is not None
    stat = {k: v for k, v in stat.items() if v["sets"]}
    for i, (naam, d) in enumerate(sorted(stat.items(),
                                         key=lambda x: (-x[1]["sets"], x[0])), start=4):
        zet(ao, i, 1, naam)
        zet(ao, i, 2, d["sets"])
        zet(ao, i, 3, len(d["dagen"]))
        zet(ao, i, 4, d["eerste"])
        zet(ao, i, 5, d["laatste"])
        zet(ao, i, 6, "{0:g} kg \u00d7 {1}".format(*d["zwaarst"]))
        zet(ao, i, 7, d["falen"])
        zet(ao, i, 8, d["rpe"])
    ao.auto_filter.ref = "A3:H{0}".format(3 + len(stat))
    ao.freeze_panes = "A4"

    # ---------- Sessies ----------
    se = wb.create_sheet("Sessies")
    zet(se, 1, 1, "Sessielog", TITEL)
    kop_rij(se, 3, ["Datum", "Titel", "Duur (min)", "Oefeningen", "Werksets",
                    "Notitie bij sessie"], [12, 30, 12, 12, 10, 50])
    for i, w in enumerate(sorted(workouts, key=lambda x: x["start_time"], reverse=True),
                          start=4):
        begin = dt.datetime.fromisoformat(w["start_time"].replace("Z", "+00:00"))
        eind = dt.datetime.fromisoformat(w["end_time"].replace("Z", "+00:00"))
        zet(se, i, 1, w["start_time"][:10])
        zet(se, i, 2, w.get("title") or "")
        zet(se, i, 3, round((eind - begin).total_seconds() / 60))
        zet(se, i, 4, len(w.get("exercises", [])))
        zet(se, i, 5, sum(len(werksets(o)) for o in w.get("exercises", [])))
        zet(se, i, 6, (w.get("description") or "")[:200])
    se.auto_filter.ref = "A3:F{0}".format(3 + len(workouts))
    se.freeze_panes = "A4"

    # ---------- Methode ----------
    me = wb.create_sheet("Methode")
    zet(me, 1, 1, "Hoe het geschatte 1RM berekend wordt", TITEL)
    zet(me, 3, 1, "Lichaamsgewicht (kg)", VET)
    zet(me, 3, 2, bw, INVOER)
    zet(me, 3, 3, "Blauw = invoerwaarde. Pas dit aan en alles rekent mee.", SUB)
    for i, regel in enumerate([
            "Dip, pull-up en muscle-up worden gescoord op het toegevoegde gewicht, maar "
            "het lichaam gaat mee omhoog.",
            "Bij de squat telt alleen het staafgewicht.",
            "",
            "Is er een RPE gelogd, dan wordt de tabel hieronder gebruikt.",
            "Is er geen RPE, dan wordt Epley gebruikt: 1RM = gewicht \u00d7 (1 + reps/30).",
            "Een set gemarkeerd als 'failure' telt als RPE 10.",
            "",
            "Sets van 7 of 8 reps krijgen een korting van 4 procent, sets van 6 twee "
            "procent, omdat extrapolatie vanaf hoge reps",
            "het maximum structureel overschat.",
            "",
            "LET OP: Dylan logde tot september 2026 vrijwel nooit RPE. De schattingen uit "
            "die periode rusten op Epley en",
            "op 'failure'-sets als RPE 10. Sinds het schema van Kaj wordt RPE wel gelogd; "
            "die regels zijn betrouwbaarder."], start=6):
        zet(me, i, 1, regel)

    tabel_rij = 20
    zet(me, tabel_rij - 1, 1, "RPE-tabel: percentage van 1RM", VET)
    kop = zet(me, tabel_rij, 1, "Reps", KOP)
    kop.fill = BLAUW
    for j, rpe in enumerate([6, 7, 8, 9, 10], start=2):
        c = zet(me, tabel_rij, j, rpe, KOP)
        c.fill = BLAUW
    for reps in range(1, 9):
        r = tabel_rij + reps
        zet(me, r, 1, reps)
        for j, rpe in enumerate([6, 7, 8, 9, 10], start=2):
            zet(me, r, j, hevy_sync.RPE_TABEL[reps][rpe], fmt="0.0%")
    for kol, breedte in zip("ABCDEF", [22, 11, 11, 11, 11, 11]):
        me.column_dimensions[kol].width = breedte

    wb.save(pad)
    return laatste - 4, len(stat), len(workouts)


def main():
    key = os.environ.get("HEVY_API_KEY", "").strip()
    if not key:
        sys.exit("HEVY_API_KEY ontbreekt.")
    pad = sys.argv[1] if len(sys.argv) > 1 else STANDAARD_PAD
    workouts = hevy_sync.alle_workouts(key)
    if not workouts:
        sys.exit("Geen workouts opgehaald.")
    bw = hevy_sync.LICHAAMSGEWICHT_TERUGVAL
    try:
        metingen = hevy_sync.api("/v1/body_measurements?page=1&pageSize=5", key)
        gewichten = [m for m in (metingen.get("body_measurements") or [])
                     if m.get("weight_kg")]
        if gewichten:
            bw = float(sorted(gewichten, key=lambda m: m.get("date", ""))[-1]["weight_kg"])
    except Exception as fout:
        print("Lichaamsgewicht niet opgehaald ({0}); {1} kg aangehouden.".format(fout, bw))
    sets, oefeningen, sessies = bouw(workouts, bw, pad)
    print("{0}\n  {1} sessies, {2} oefeningen, {3} werksets op de wedstrijdlifts, "
          "lichaamsgewicht {4} kg".format(pad, sessies, oefeningen, sets, bw))


if __name__ == "__main__":
    main()
